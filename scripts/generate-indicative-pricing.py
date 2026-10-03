"""Build the sanitized Preview estimate table from the approved research workbook.

This is intentionally an internal build step. It never exports supplier costs,
quantities, margins, or supplier/depot data to the browser bundle.
"""

from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = Path(
    r"C:\Users\RJ Gadgetz\Documents\Codex\2026-07-24\realtime-voice-chat\pricing-handoff\Container_Exchange_Market_Pricing_US_Canada_v3.xlsx"
)
PUBLIC_OUT = ROOT / "src" / "data" / "indicativeMarketPrices.js"
INTERNAL_OUT = ROOT / "pricing-research" / "indicative-pricing-assessment.json"


def slug(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "-", str(value or "").lower()).strip("-")


def product_id(size: str, grade: str) -> str | None:
    size = str(size).lower()
    grade = str(grade).lower()
    length = "20" if "20ft" in size else "40" if "40ft" in size else None
    if not length:
        return None
    high_cube = "high cube" in size
    suffix = "hc" if high_cube else ""
    if grade == "one trip":
        return f"new-{length}{suffix}-iicl"
    grade_map = {
        "iicl": "iicl",
        "cargo worthy": "cw",
        "wind & water tight": "wwt",
        "as-is": "as-is",
    }
    mapped = grade_map.get(grade)
    return f"used-{length}{suffix}-{mapped}" if mapped else None


def market_id(city: str, state: str) -> str:
    special = {
        ("Los Angeles", "CA"): "los-angeles-long-beach-ca",
        ("Long Beach", "CA"): "los-angeles-long-beach-ca",
        ("New York", "NY"): "new-york-newark-ny-nj",
        ("Newark", "NJ"): "new-york-newark-ny-nj",
        ("San Francisco", "CA"): "san-francisco-oakland-ca",
        ("Oakland", "CA"): "san-francisco-oakland-ca",
        ("Boston", "MA"): "worcester-boston-ma",
        ("Worcester", "MA"): "worcester-boston-ma",
        ("Vancouver", "BC"): "vancouver-delta-bc",
        ("Delta", "BC"): "vancouver-delta-bc",
        ("Halifax", "NS"): "halifax-dartmouth-ns",
        ("Dartmouth", "NS"): "halifax-dartmouth-ns",
    }
    return special.get((str(city), str(state)), f"{slug(city)}-{slug(state)}")


def main() -> None:
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    sheet = workbook["Price Recommendations"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=4, max_row=4))]
    rows = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=5, values_only=True) if values[0]]

    public: dict[str, dict[str, dict]] = {}
    assessment = []
    reasons = Counter()
    estimate_count = 0
    below_profit_target = 0

    for row in rows:
        country = str(row.get("Country") or "")
        state = str(row.get("State / Province") or "")
        city = str(row.get("City") or "")
        currency = str(row.get("Currency") or "")
        evidence_status = str(row.get("Evidence Status") or "")
        proposed = row.get("Proposed Container Exchange price")
        product = product_id(row.get("Container size"), row.get("Container grade"))
        source_url = str(row.get("Benchmark source URL") or "").strip()
        source_date = row.get("Benchmark date checked")
        confidence = str(row.get("Confidence") or "")
        explanation = str(row.get("Derived / estimate explanation and review note") or "").strip()
        market = market_id(city, state)

        eligible = True
        reason = ""
        if country != "US":
            eligible, reason = False, "Canadian pricing remains on currency hold"
        elif currency != "USD":
            eligible, reason = False, "Currency is not confirmed as USD"
        elif not product:
            eligible, reason = False, "No exact catalog product mapping"
        elif not isinstance(proposed, (int, float)) or proposed <= 0:
            eligible, reason = False, "No usable proposed price"
        elif not source_url or not source_date:
            eligible, reason = False, "Missing benchmark source or check date"
        elif evidence_status == "Source does not substantiate price":
            eligible, reason = False, "Source does not substantiate price"

        profit = row.get("Estimated gross profit")
        below_target = isinstance(profit, (int, float)) and profit < 700
        if below_target:
            below_profit_target += 1

        if eligible:
            record = {
                "country": "US",
                "currency": "USD",
                "price": int(proposed),
            }
            public.setdefault(market, {})[product] = record
            estimate_count += 1
        else:
            reasons[reason] += 1

        assessment.append(
            {
                "country": country,
                "state": state,
                "city": city,
                "marketId": market,
                "productId": product,
                "currency": currency,
                "proposedPrice": proposed,
                "assessment": "indicative_estimate" if eligible else "request_quote",
                "reason": reason or None,
                "sourceDate": str(source_date) if source_date else None,
                "evidenceStatus": evidence_status,
                "confidence": confidence,
                "explanation": explanation,
                "belowFixedPriceProfitTarget": below_target,
            }
        )

    # The client bundle deliberately contains only the consumer-facing amount,
    # currency and exact market/product mapping.
    PUBLIC_OUT.write_text(
        "// Generated from the internal market research workbook. Do not edit by hand.\n"
        "// Estimates are not fixed prices and never authorize checkout.\n"
        f"export const INDICATIVE_MARKET_PRICES = {json.dumps(public, indent=2, sort_keys=True)};\n",
        encoding="utf-8",
    )
    INTERNAL_OUT.parent.mkdir(parents=True, exist_ok=True)
    INTERNAL_OUT.write_text(
        json.dumps(
            {
                "sourceWorkbook": WORKBOOK.name,
                "rowsAssessed": len(rows),
                "indicativeEstimateRows": estimate_count,
                "quoteOnlyRows": len(rows) - estimate_count,
                "quoteOnlyReasons": reasons,
                "belowFixedPriceProfitTarget": below_profit_target,
                "rows": assessment,
            },
            indent=2,
            default=str,
        ),
        encoding="utf-8",
    )
    print(json.dumps({"estimates": estimate_count, "quoteOnly": len(rows) - estimate_count, "reasons": reasons, "belowProfitTarget": below_profit_target}, default=str))


if __name__ == "__main__":
    main()
